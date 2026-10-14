from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import joinedload

from backend.app.db.database import session_local
from backend.app.db.models import CompanyMaster, CustomerMaster, InvoiceHeader, ProductCatalog

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "data" / "us_invoice_management_data.xlsx"
HEADER_FILL = PatternFill("solid", fgColor="173B69")


def add_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        sheet.column_dimensions[column[0].column_letter].width = width


def main() -> None:
    db = session_local()
    try:
        companies = db.scalars(select(CompanyMaster).order_by(CompanyMaster.company_name)).all()
        customers = db.scalars(select(CustomerMaster).order_by(CustomerMaster.customer_name)).all()
        products = db.scalars(select(ProductCatalog).order_by(ProductCatalog.product_name)).all()
        invoices = db.scalars(select(InvoiceHeader).options(joinedload(InvoiceHeader.items)).order_by(InvoiceHeader.invoice_no)).unique().all()

        workbook = Workbook()
        workbook.remove(workbook.active)
        add_sheet(workbook, "Companies", ["Company ID", "Company Name", "Phone", "Email", "Address", "Ship From Address"], [
            [item.company_id, item.company_name, item.phone, item.email, item.address, item.ship_from_address] for item in companies
        ])
        add_sheet(workbook, "Customers", ["Customer ID", "Customer Name", "Contact Person", "Phone", "Email", "Billing Address", "Shipping Address", "State / Region"], [
            [item.customer_id, item.customer_name, item.contact_person, item.phone, item.email, item.billing_address, item.shipping_address, item.state_code] for item in customers
        ])
        add_sheet(workbook, "Products", ["Product ID", "Product Name", "HSN/SAC", "Unit", "Rate USD", "Tax %"], [
            [item.product_id, item.product_name, item.hsn_sac, item.unit, item.rate, item.gst_percentage] for item in products
        ])
        add_sheet(workbook, "Invoices", ["Invoice Number", "Invoice Date", "Due Date", "PO Reference", "Company ID", "Customer ID", "Place of Supply", "Status", "Currency", "Subtotal USD", "Tax USD", "Grand Total USD", "Amount in Words"], [
            [item.invoice_no, item.invoice_date, item.due_date, item.po_reference, item.company_id, item.customer_id, item.place_of_supply, item.invoice_status, "USD", item.subtotal, item.igst + item.cgst + item.sgst, item.grand_total, item.amount_in_words] for item in invoices
        ])
        add_sheet(workbook, "Invoice Line Items", ["Invoice Number", "Product ID", "Description", "HSN/SAC", "Quantity", "Unit", "Unit Price USD", "Tax %", "Line Amount USD"], [
            [invoice.invoice_no, line.product_id, line.description, line.hsn_sac, line.quantity, line.unit, line.rate, line.gst_percentage, line.amount]
            for invoice in invoices for line in invoice.items
        ])
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(OUTPUT)
        print(f"Created {OUTPUT}")
    finally:
        db.close()


if __name__ == "__main__":
    main()
